"""Excel workbook input parser."""

from pathlib import Path

from openpyxl import load_workbook

from parsers.base import DataParser
from parsers.csv_parser import CsvParser


class ExcelParser(DataParser):
    def parse(self, file_path: Path) -> list[dict]:
        workbook = load_workbook(file_path, data_only=True, read_only=True)
        return self._sheet_to_records(workbook[workbook.sheetnames[0]])

    def parse_all_sheets(self, file_path: Path) -> dict[str, list[dict]]:
        workbook = load_workbook(file_path, data_only=True, read_only=True)
        return {name: self._sheet_to_records(workbook[name]) for name in workbook.sheetnames}

    @staticmethod
    def _sheet_to_records(sheet) -> list[dict]:
        rows = sheet.iter_rows(values_only=True)
        try:
            first_row = next(rows)
        except StopIteration:
            return []

        headers = [str(value).strip() if value is not None else f"col_{index}"
                   for index, value in enumerate(first_row)]
        records = []
        for row in rows:
            if all(value is None for value in row):
                continue
            record = dict(zip(headers, row))
            records.append(CsvParser._decode_nested_cells(record))
        return records
